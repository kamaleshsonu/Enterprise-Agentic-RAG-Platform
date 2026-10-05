from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from bs4 import BeautifulSoup


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".html",
    ".txt",
}


def extract_pdf(file_path: Path) -> str:
    reader = PdfReader(file_path)

    text = []

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text.append(page_text)

    return "\n".join(text)


def extract_docx(file_path: Path) -> str:
    document = Document(file_path)

    text = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    return "\n".join(text)


def extract_pptx(file_path: Path) -> str:
    presentation = Presentation(file_path)

    text = []

    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text") and shape.text.strip():
                text.append(shape.text)

    return "\n".join(text)


def extract_html(file_path: Path) -> str:
    html = file_path.read_text(
        encoding="utf-8",
        errors="ignore"
    )

    soup = BeautifulSoup(html, "html.parser")

    return soup.get_text(
        separator="\n",
        strip=True
    )


def extract_txt(file_path: Path) -> str:
    return file_path.read_text(
        encoding="utf-8",
        errors="ignore"
    )


def extract_text(file_path: Path) -> str:

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".pptx":
        return extract_pptx(file_path)

    if extension == ".html":
        return extract_html(file_path)

    if extension == ".txt":
        return extract_txt(file_path)

    return ""


def scan_data(data_path: str):

    data_dir = Path(data_path)

    if not data_dir.exists():
        print(f"ERROR: Data directory not found: {data_dir}")
        return

    print(f"\nScanning: {data_dir.resolve()}\n")

    for dataset_dir in data_dir.iterdir():

        if not dataset_dir.is_dir():
            continue

        dataset_name = dataset_dir.name

        print(f"========== {dataset_name.upper()} ==========")

        for file_path in dataset_dir.rglob("*"):

            if not file_path.is_file():
                continue

            extension = file_path.suffix.lower()

            if extension not in SUPPORTED_EXTENSIONS:
                continue

            text = extract_text(file_path)

            print(f"\nFILE: {file_path.name}")
            print(f"TYPE: {extension}")
            print(f"CHARACTERS EXTRACTED: {len(text)}")


if __name__ == "__main__":
    scan_data("DATA")